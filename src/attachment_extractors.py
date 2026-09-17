from pathlib import Path
import csv
import io
import json
import subprocess
import tempfile

import pymupdf
import pytesseract
from PIL import Image
from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook
import xlrd
from pptx import Presentation
from striprtf.striprtf import rtf_to_text
from bs4 import BeautifulSoup
from faster_whisper import WhisperModel


TESSERACT_EXE = r"C:\Program Files\Tesseract-OCR\tesseract.exe"
FFMPEG_EXE = r"C:\ffmpeg\bin\ffmpeg.exe"

pytesseract.pytesseract.tesseract_cmd = TESSERACT_EXE


IMAGE_TYPES = {
    "image/png",
    "image/jpeg",
    "image/jpg",
}

AUDIO_TYPES = {
    "audio/mpeg",
    "audio/mp3",
    "audio/wav",
    "audio/x-wav",
    "audio/wave",
    "audio/mp4",
    "audio/x-m4a",
    "audio/ogg",
    "audio/webm",
}

VIDEO_TYPES = {
    "video/mp4",
    "video/webm",
    "video/x-msvideo",
    "video/quicktime",
    "video/x-matroska",
}


TEXT_EXTENSIONS = {
    ".txt",
    ".md",
    ".log",
    ".ini",
    ".cfg",
    ".conf",
    ".xml",
    ".yaml",
    ".yml",
    ".sql",
    ".py",
    ".js",
    ".jsx",
    ".ts",
    ".tsx",
    ".css",
    ".java",
    ".c",
    ".cpp",
    ".h",
    ".hpp",
}


_whisper_model = None


def get_whisper_model():
    global _whisper_model

    if _whisper_model is None:
        _whisper_model = WhisperModel(
            "small",
            device="cpu",
            compute_type="int8"
        )

    return _whisper_model


def extract_image_text(file_path):
    path = Path(file_path)

    with Image.open(path) as image:
        text = pytesseract.image_to_string(image)

    return text.strip()


def extract_pdf_text(file_path):
    """
    Extract text from normal text-based PDF pages.

    If a page contains no extractable text, render that page
    and run Tesseract OCR on it.

    This means mixed PDFs are also supported:
    text extraction for normal pages + OCR for scanned pages.
    """

    path = Path(file_path)
    pages = []

    try:
        reader = PdfReader(str(path))

        for index, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""

            if text.strip():
                pages.append(
                    f"[Page {index}]\n{text.strip()}"
                )
                continue

            try:
                with pymupdf.open(str(path)) as pdf:
                    pdf_page = pdf[index - 1]
                    pixmap = pdf_page.get_pixmap(
                        matrix=pymupdf.Matrix(2, 2),
                        alpha=False
                    )

                    image_bytes = pixmap.tobytes("png")

                    with Image.open(
                        io.BytesIO(image_bytes)
                    ) as image:
                        ocr_text = pytesseract.image_to_string(
                            image
                        ).strip()

                    if ocr_text:
                        pages.append(
                            f"[Page {index}]\n{ocr_text}"
                        )

            except Exception:
                continue

    except Exception:
        pages = []

    return "\n\n".join(pages).strip()


def extract_docx_text(file_path):
    document = Document(str(file_path))
    parts = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            parts.append(text)

    for table in document.tables:
        for row in table.rows:
            values = [
                cell.text.strip()
                for cell in row.cells
            ]

            if any(values):
                parts.append(" | ".join(values))

    for section in document.sections:
        for paragraph in section.header.paragraphs:
            text = paragraph.text.strip()

            if text:
                parts.append(text)

        for table in section.header.tables:
            for row in table.rows:
                values = [
                    cell.text.strip()
                    for cell in row.cells
                ]

                if any(values):
                    parts.append(" | ".join(values))

        for paragraph in section.footer.paragraphs:
            text = paragraph.text.strip()

            if text:
                parts.append(text)

        for table in section.footer.tables:
            for row in table.rows:
                values = [
                    cell.text.strip()
                    for cell in row.cells
                ]

                if any(values):
                    parts.append(" | ".join(values))

    return "\n".join(parts).strip()


def extract_xlsx_text(file_path):
    workbook = load_workbook(
        filename=str(file_path),
        read_only=True,
        data_only=False
    )

    parts = []

    for sheet in workbook.worksheets:
        parts.append(f"[Sheet: {sheet.title}]")

        for row in sheet.iter_rows(values_only=True):
            values = []

            for value in row:
                if value is not None:
                    values.append(str(value).strip())

            if values:
                parts.append(" | ".join(values))

    workbook.close()

    return "\n".join(parts).strip()


def extract_xls_text(file_path):
    workbook = xlrd.open_workbook(str(file_path))
    parts = []

    for sheet in workbook.sheets():
        parts.append(f"[Sheet: {sheet.name}]")

        for row_index in range(sheet.nrows):
            values = []

            for col_index in range(sheet.ncols):
                value = sheet.cell_value(
                    row_index,
                    col_index
                )

                if value != "":
                    values.append(str(value).strip())

            if values:
                parts.append(" | ".join(values))

    return "\n".join(parts).strip()


def extract_csv_text(file_path):
    parts = []

    with open(
        file_path,
        "r",
        encoding="utf-8-sig",
        errors="replace",
        newline=""
    ) as file:
        reader = csv.reader(file)

        for row in reader:
            values = [
                value.strip()
                for value in row
                if value.strip()
            ]

            if values:
                parts.append(" | ".join(values))

    return "\n".join(parts).strip()


def extract_pptx_text(file_path):
    presentation = Presentation(str(file_path))
    parts = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1
    ):
        slide_parts = []

        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text = shape.text.strip()

                if text:
                    slide_parts.append(text)

            if getattr(shape, "has_table", False):
                table = shape.table

                for row in table.rows:
                    values = [
                        cell.text.strip()
                        for cell in row.cells
                    ]

                    if any(values):
                        slide_parts.append(
                            " | ".join(values)
                        )

        try:
            notes_slide = slide.notes_slide

            for shape in notes_slide.shapes:
                if hasattr(shape, "text"):
                    text = shape.text.strip()

                    if text:
                        slide_parts.append(text)
        except Exception:
            pass

        if slide_parts:
            parts.append(
                f"[Slide {slide_number}]\n"
                + "\n".join(slide_parts)
            )

    return "\n\n".join(parts).strip()


def extract_text_file(file_path):
    path = Path(file_path)

    raw = path.read_bytes()

    encodings = [
        "utf-8-sig",
        "utf-8",
        "utf-16",
        "utf-16-le",
        "utf-16-be",
        "cp1252",
        "latin-1",
    ]

    for encoding in encodings:
        try:
            return raw.decode(encoding).strip()
        except UnicodeDecodeError:
            continue

    return raw.decode(
        "utf-8",
        errors="replace"
    ).strip()


def extract_html_text(file_path):
    raw = Path(file_path).read_bytes()

    html = raw.decode(
        "utf-8",
        errors="replace"
    )

    soup = BeautifulSoup(
        html,
        "html.parser"
    )

    for element in soup(
        ["script", "style", "noscript"]
    ):
        element.decompose()

    parts = []

    if soup.title and soup.title.string:
        title = soup.title.string.strip()

        if title:
            parts.append(f"[Title]\n{title}")

    body_text = soup.get_text(
        "\n",
        strip=True
    )

    if body_text:
        parts.append(body_text)

    return "\n\n".join(parts).strip()


def extract_json_text(file_path):
    raw = extract_text_file(file_path)

    try:
        data = json.loads(raw)

        return json.dumps(
            data,
            ensure_ascii=False,
            indent=2
        ).strip()

    except Exception:
        return raw.strip()


def extract_rtf_text(file_path):
    raw = extract_text_file(file_path)

    return rtf_to_text(raw).strip()


def extract_audio_text(file_path):
    model = get_whisper_model()

    segments, _ = model.transcribe(
        str(Path(file_path).resolve()),
        beam_size=5,
        vad_filter=True
    )

    parts = []

    for segment in segments:
        text = segment.text.strip()

        if text:
            parts.append(text)

    return " ".join(parts).strip()


def extract_video_audio_text(file_path):
    temp_path = None

    try:
        with tempfile.NamedTemporaryFile(
            suffix=".wav",
            delete=False
        ) as temp_file:
            temp_path = temp_file.name

        subprocess.run(
            [
                FFMPEG_EXE,
                "-i",
                str(Path(file_path).resolve()),
                "-vn",
                "-ar",
                "16000",
                "-ac",
                "1",
                "-c:a",
                "pcm_s16le",
                temp_path,
                "-y",
                "-loglevel",
                "error"
            ],
            check=True
        )

        return extract_audio_text(temp_path)

    finally:
        if temp_path:
            Path(temp_path).unlink(
                missing_ok=True
            )


def extract_attachment_text(file_path, mime_type):
    path = Path(file_path)
    extension = path.suffix.lower()

    if mime_type in IMAGE_TYPES:
        text = extract_image_text(path)

        if text:
            return text, "extracted"

        return None, "no_text"

    if (
        mime_type == "application/pdf"
        or extension == ".pdf"
    ):
        text = extract_pdf_text(path)

        if text:
            return text, "extracted"

        return None, "no_text"

    if (
        mime_type
        == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        or extension == ".docx"
    ):
        text = extract_docx_text(path)

        if text:
            return text, "extracted"

        return None, "no_text"

    if (
        mime_type
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        or extension == ".xlsx"
    ):
        text = extract_xlsx_text(path)

        if text:
            return text, "extracted"

        return None, "no_text"

    if (
        mime_type == "application/vnd.ms-excel"
        or extension == ".xls"
    ):
        text = extract_xls_text(path)

        if text:
            return text, "extracted"

        return None, "no_text"

    if (
        mime_type == "text/csv"
        or extension == ".csv"
    ):
        text = extract_csv_text(path)

        if text:
            return text, "extracted"

        return None, "no_text"

    if (
        mime_type
        == "application/vnd.openxmlformats-officedocument.presentationml.presentation"
        or extension == ".pptx"
    ):
        text = extract_pptx_text(path)

        if text:
            return text, "extracted"

        return None, "no_text"

    if (
        mime_type == "text/html"
        or extension in {".html", ".htm"}
    ):
        text = extract_html_text(path)

        if text:
            return text, "extracted"

        return None, "no_text"

    if (
        mime_type == "application/json"
        or extension == ".json"
    ):
        text = extract_json_text(path)

        if text:
            return text, "extracted"

        return None, "no_text"

    if (
        mime_type == "application/rtf"
        or mime_type == "text/rtf"
        or extension == ".rtf"
    ):
        text = extract_rtf_text(path)

        if text:
            return text, "extracted"

        return None, "no_text"

    if extension in TEXT_EXTENSIONS:
        text = extract_text_file(path)

        if text:
            return text, "extracted"

        return None, "no_text"

    if (
        mime_type in AUDIO_TYPES
        or extension in {
            ".mp3",
            ".wav",
            ".m4a",
            ".ogg",
            ".webm"
        }
    ):
        text = extract_audio_text(path)

        if text:
            return text, "extracted"

        return None, "no_text"

    if (
        mime_type in VIDEO_TYPES
        or extension in {
            ".mp4",
            ".mkv",
            ".avi",
            ".mov",
            ".webm"
        }
    ):
        text = extract_video_audio_text(path)

        if text:
            return text, "extracted"

        return None, "no_text"

    return None, "unsupported"
